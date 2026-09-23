from __future__ import annotations

import json
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "社區故事" / "已確認"
OUTPUT = ROOT / "查核資料" / "已確認故事擷取.json"


def read_doc(path: Path) -> dict:
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)
    return {
        "folder": path.parent.name,
        "file": path.name,
        "paragraphs": paragraphs,
        "tables": tables,
    }


def main() -> None:
    records = [read_doc(path) for path in sorted(SOURCE.rglob("*.docx"))]
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {OUTPUT} with {len(records)} documents")


if __name__ == "__main__":
    main()
